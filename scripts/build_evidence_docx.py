#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""JSON -> 佐证材料汇编 docx（编号 + 一行说明 + 居中图片）。

用法:
    python3 build_evidence_docx.py config.json

配置示例:
{
  "out": "/tmp/佐证材料汇编.docx",
  "header": {
    "org": "<学院>",
    "title": "综合测评佐证材料",
    "meta": "姓名：<姓名>　　学号：<学号>　　班级：<班级>",
    "note": "（以下材料按《…明细表（个人提交）》填报顺序编号）"
  },
  "prepare_dir": "/tmp/evidence_img",
  "entries": [
    {"caption": "【1】科学文化素质·发展分——参加学术讲座：<讲座名>（<日期>，本人现场拍摄）",
     "images": ["/abs/shot.png"]},
    {"caption": "【2】劳动技能素质·基础分——<任职/称号>（公示名单第 22 号）",
     "images": ["/abs/公示.pdf#page=2"]}
  ]
}

图片写法:
    /path/a.jpg            普通图片
    /path/公示.pdf#page=2  PDF 指定页渲染（公示页、秩序册页首选）
"""
from __future__ import annotations

import argparse
import json
import os
import sys

try:
    from PIL import Image, ImageOps
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
except ImportError as e:  # noqa: BLE001
    print(f"缺少依赖 ({e})：pip3 install python-docx pillow pymupdf", file=sys.stderr)
    raise SystemExit(2)


def cjk(run, font):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)


def prepare(src: str, out_dir: str, name: str, dpi: int = 150, max_px: int = 1600, quality: int = 82):
    """返回 (图片路径, (w, h))。支持 `x.pdf#page=N`。"""
    os.makedirs(out_dir, exist_ok=True)
    dst = os.path.join(out_dir, name + ".jpg")
    if ".pdf#page=" in src:
        import pymupdf

        path, page = src.split("#page=")
        pix = pymupdf.open(path)[int(page) - 1].get_pixmap(dpi=dpi)
        pix.save(dst)
        with Image.open(dst) as im:
            size = im.size
            im.convert("RGB").save(dst, quality=quality, optimize=True)
        return dst, size

    im = ImageOps.exif_transpose(Image.open(src)).convert("RGB")
    im.thumbnail((max_px, max_px))
    im.save(dst, quality=quality, optimize=True)
    return dst, im.size


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("config")
    args = ap.parse_args()

    with open(args.config, encoding="utf-8") as f:
        cfg = json.load(f)

    out_dir = cfg.get("prepare_dir", "/tmp/zongce_evidence_img")
    head = cfg.get("header", {})
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.2)
    sec.left_margin = sec.right_margin = Cm(2.2)

    def para(text, size=12, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, font="黑体", space_after=6):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.space_before = Pt(0)
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.bold = bold
        cjk(r, font)
        return p

    if head.get("org"):
        para(head["org"], size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font="黑体", space_after=0)
    if head.get("title"):
        para(head["title"], size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font="黑体", space_after=6)
    if head.get("meta"):
        para(head["meta"], size=11, align=WD_ALIGN_PARAGRAPH.CENTER, font="宋体", space_after=2)
    if head.get("note"):
        para(head["note"], size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, font="宋体", space_after=14)

    w_land = Cm(cfg.get("width_landscape_cm", 15))
    w_port = Cm(cfg.get("width_portrait_cm", 11))
    total = 0
    for idx, entry in enumerate(cfg["entries"], start=1):
        caption = entry.get("caption", "")
        if caption:
            para(caption, size=12, bold=True, font="黑体", space_after=4)
        for j, src in enumerate(entry.get("images", []), start=1):
            name = f"{idx:02d}_{j}"
            path, (w, h) = prepare(src, out_dir, name, dpi=cfg.get("pdf_dpi", 150),
                                   max_px=cfg.get("max_px", 1600), quality=cfg.get("jpeg_quality", 82))
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(10)
            p.add_run().add_picture(path, width=(w_land if w >= h else w_port))
            total += 1
        if not caption and not entry.get("images"):
            print(f"warning: entry {idx} 既无说明也无图片", file=sys.stderr)

    doc.save(cfg["out"])
    print(f"saved: {cfg['out']}  entries {len(cfg['entries'])}  images {total}  "
          f"{os.path.getsize(cfg['out']) / 1024 / 1024:.2f} MB")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
