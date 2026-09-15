#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""JSON -> 综合测评明细表 xlsx（类别/项目/加分分值，自动合并类别列）。

用法:
    python3 build_detail_xlsx.py config.json

配置示例（demo.json）:
{
  "out": "/tmp/明细表.xlsx",
  "title": "<学院>综合测评细则明细表",
  "meta": {"班级": "<班级>", "姓名": "<姓名>", "学号": "<学号>"},
  "headers": ["所属类别", "项目", "加分分值（100分制）"],
  "rows": [
    {"cat": "思想道德素质评价", "sub": "发展分", "item": "（无）", "score": 0},
    {"cat": "科学文化素质评价", "sub": "发展分", "item": "参加学术讲座：<讲座A>", "score": null},
    {"item": "参加学术讲座：<讲座B>", "score": null},
    {"item": "", "score": 5},
    {"cat": "身心健康素质评价", "sub": "发展分", "item": "<趣味项目>（趣味减半）", "score": 0.5}
  ],
  "signature": "学生签字（手写）：",
  "note": "注：请出具相关的证明（证书/文章的复印件、作品原件），并按照上述顺序依次排序。"
}

约定:
    - 省略 cat/sub 表示继承上一行（用于同一类别下的多行项目）；
    - item 为空字符串 -> 该行项目格留空；
    - score 为 null -> 分值格留空；写 0 会显示 0（"按次计分"项未达阈值时就该写 0）；
    - 每行可用 "height" 指定行高，否则按 C 列文本长度自动估算（避免长项目名被截断）；
    - A 列按"连续相同 cat"合并，B 列按"连续相同 cat+sub"合并。
"""
from __future__ import annotations

import argparse
import json
import sys

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, Side
except ImportError:
    print("需要 openpyxl: pip3 install openpyxl", file=sys.stderr)
    raise SystemExit(2)

DEFAULTS = {
    "title_font": "方正小标宋_GBK",
    "head_font": "黑体",
    "cat_font": "黑体",
    "body_font": "宋体",
    "title_size": 20,
    "size": 11,
    "row_height": 20,
    "col_widths": [21.625, 12, 51.625, 33.375],
    "note_height": 29,
}


def est_height(text, col_width, base=20.0, line_unit=15.0):
    """按 C 列文本的视觉宽度估算行高：行数 × 单行文字高度，且不低于 base。

    列宽单位 ≈ 1 个半角字符宽；全角字符按 2 计，ASCII 按 1 计。
    单行 -> 20（模板默认），两行 -> 30（≈两倍行高），三行 -> 45。
    """
    if not text:
        return base
    vis = sum(2 if ord(ch) > 0x2E80 else 1 for ch in str(text))
    w = max(int(col_width), 1)
    lines = (vis + w - 1) // w
    return max(base, lines * line_unit)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("config")
    args = ap.parse_args()

    with open(args.config, encoding="utf-8") as f:
        cfg = json.load(f)

    st = dict(DEFAULTS)
    st.update(cfg.get("style", {}))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = cfg.get("sheet", "Sheet1")

    thin = Side(style="thin")
    bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    right = Alignment(horizontal="right", vertical="center")

    f_title = Font(name=st["title_font"], size=st["title_size"], bold=True)
    f_head = Font(name=st["head_font"], size=st["size"], bold=True)
    f_cat = Font(name=st["cat_font"], size=st["size"], bold=True)
    f_song = Font(name=st["body_font"], size=st["size"])

    headers = cfg.get("headers", ["所属类别", "项目", "加分分值（100分制）"])
    meta = cfg.get("meta", {})
    rows = cfg["rows"]

    # ---- title ----
    ws.merge_cells("A1:D3")
    ws["A1"] = cfg.get("title", "综合测评明细表")
    ws["A1"].font = f_title
    ws["A1"].alignment = center

    # ---- meta line ----
    meta_items = list(meta.items())
    keys = [f"{k}：{v}" for k, v in meta_items]
    ws.merge_cells("A4:B5")
    ws["A4"] = keys[0] if keys else ""
    if len(keys) > 1:
        ws.merge_cells("C4:C5")
        ws["C4"] = keys[1]
    if len(keys) > 2:
        ws.merge_cells("D4:D5")
        ws["D4"] = keys[2]
    for c in ("A4", "C4", "D4"):
        ws[c].font = f_song
        ws[c].alignment = left

    # ---- header ----
    ws.merge_cells("A6:B6")
    ws["A6"] = headers[0]
    ws["C6"] = headers[1]
    ws["D6"] = headers[2]
    for c in ("A6", "B6", "C6", "D6"):
        ws[c].font = f_head
        ws[c].alignment = center
        ws[c].border = bd

    # ---- body ----
    start = 7
    last_cat = ""
    last_sub = ""
    for i, row in enumerate(rows):
        r = start + i
        cat = row.get("cat", last_cat)
        sub = row.get("sub", last_sub)
        last_cat, last_sub = cat, sub
        row["cat"], row["sub"] = cat, sub  # 写回继承值，供后续合并推导使用
        if cat:
            ws.cell(row=r, column=1, value=cat)
        if sub:
            ws.cell(row=r, column=2, value=sub)
        item = row.get("item", "")
        if item:
            ws.cell(row=r, column=3, value=item)
        score = row.get("score")
        if score is not None:
            ws.cell(row=r, column=4, value=score)

    end = start + len(rows) - 1

    # 合并连续相同 cat 的 A 列、连续相同 (cat, sub) 的 B 列
    def merge_col(col: int, same_key):
        run_start = start
        key = same_key(_row_at(start))
        for r in range(start + 1, end + 2):
            cur = same_key(_row_at(r)) if r <= end else None
            if cur != key:
                if r - 1 > run_start:
                    ws.merge_cells(start_row=run_start, start_column=col,
                                   end_row=r - 1, end_column=col)
                run_start = r
                key = cur

    def _row_at(r):
        if r < start or r > end:
            return {"cat": None, "sub": None}
        i = r - start
        return rows[i]

    merge_col(1, lambda x: x.get("cat"))
    merge_col(2, lambda x: (x.get("cat"), x.get("sub")))

    for r in range(start, end + 1):
        ws.row_dimensions[r].height = rows[r - start].get("height") or est_height(
            rows[r - start].get("item", ""), st["col_widths"][2], st["row_height"])
        for col in range(1, 5):
            c = ws.cell(row=r, column=col)
            c.border = bd
            if col <= 2:
                c.font = f_cat
                c.alignment = center
            elif col == 3:
                c.font = f_song
                c.alignment = left
            else:
                c.font = f_song
                c.alignment = center

    # ---- signature + note ----
    sig = start + len(rows) + 1
    ws.merge_cells(start_row=sig, start_column=1, end_row=sig + 1, end_column=3)
    ws.cell(row=sig, column=1, value=cfg.get("signature", "学生签字（手写）："))
    ws.merge_cells(start_row=sig, start_column=4, end_row=sig + 1, end_column=4)
    for r in (sig, sig + 1):
        ws.row_dimensions[r].height = st["row_height"]
        for col in range(1, 5):
            c = ws.cell(row=r, column=col)
            c.border = bd
            c.font = f_song
            c.alignment = right

    note_row = sig + 2
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=4)
    camel = ws.cell(row=note_row, column=1, value=cfg.get(
        "note", "注：请出具相关的证明（证书/文章的复印件、作品原件），并按照上述顺序依次排序。"))
    camel.font = Font(name=st["head_font"], size=st["size"], bold=True)
    camel.alignment = left
    ws.row_dimensions[note_row].height = st["note_height"]

    widths = st["col_widths"]
    for idx, w in enumerate(widths):
        ws.column_dimensions[chr(ord("A") + idx)].width = w

    wb.save(cfg["out"])
    print(f"saved: {cfg['out']}  rows {start}-{end}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
